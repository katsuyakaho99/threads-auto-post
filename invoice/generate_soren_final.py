"""株式会社蒼蓮あて 請求書（提出用・清書版）を生成する。
記入用の黄色セル・凡例・備考は付けず、体裁を整えた最終版。"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE = os.path.dirname(os.path.abspath(__file__))
FONT = "Arial"
NAVY = "1F3864"
BAND = "EAEFF7"
GRAY = "BFBFBF"

navy_fill = PatternFill("solid", fgColor=NAVY)
band_fill = PatternFill("solid", fgColor=BAND)
thin = Side(style="thin", color=GRAY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
navy_side = Side(style="medium", color=NAVY)


def build(path, inv_no, issue_date, due_date, items):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "請求書"
    ws.sheet_view.showGridLines = False

    def c(coord, value=None, size=11, bold=False, color="000000", align="left",
          valign="center", fill=None, border=None, wrap=False, fmt=None):
        cell = ws[coord]
        if value is not None:
            cell.value = value
        cell.font = Font(name=FONT, size=size, bold=bold, color=color)
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        if fill:
            cell.fill = fill
        if border:
            cell.border = border
        if fmt:
            cell.number_format = fmt
        return cell

    widths = {"A": 5, "B": 42, "C": 7, "D": 14, "E": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # ---- タイトル ----
    ws.row_dimensions[1].height = 8
    ws.merge_cells("A2:E2")
    c("A2", "請求書", size=26, bold=True, color=NAVY, align="center")
    ws.row_dimensions[2].height = 40
    for col in "ABCDE":
        ws[f"{col}3"].border = Border(bottom=navy_side)
    ws.row_dimensions[3].height = 4

    # ---- 宛先（左）----
    ws.merge_cells("A5:C5")
    c("A5", "株式会社蒼蓮　御中", size=15, bold=True)
    ws[f"A5"].border = Border(bottom=Side(style="thin", color="808080"))
    for col in "BC":
        ws[f"{col}5"].border = Border(bottom=Side(style="thin", color="808080"))
    c("A6", "ご担当：小笹彩加 様", size=10)
    ws.merge_cells("A7:C7")
    c("A7", "〒542-0012 大阪市中央区谷町6丁目3番25号 LOK09号", size=9, color="404040")

    # ---- 請求メタ（右）----
    c("D5", "請求日", size=10, bold=True, align="right")
    c("E5", issue_date, size=10, align="right")
    c("D6", "請求書番号", size=10, bold=True, align="right")
    c("E6", inv_no, size=10, align="right")
    c("D7", "支払期限", size=10, bold=True, align="right")
    c("E7", due_date, size=10, align="right")

    # ---- ご請求金額 ----
    ws.merge_cells("A9:B9")
    c("A9", "ご請求金額（税込）", size=12, bold=True, color="FFFFFF", align="center", fill=navy_fill)
    ws.merge_cells("C9:E9")
    total_cell = None  # set later after we know total row

    # ---- 発行者（右）----
    c("A10", "下記のとおりご請求申し上げます。", size=10)
    ws.merge_cells("D10:E10")
    c("D10", "勝矢 夏帆", size=12, bold=True, align="right")
    ws.merge_cells("D11:E12")
    c("D11", "〒550-0003 大阪府大阪市西区京町堀2-9-18\nジュネーゼ京町堀パークサイド601",
      size=9, color="404040", align="right", valign="top", wrap=True)

    # ---- 明細テーブル ----
    header_row = 14
    heads = [("A", "No."), ("B", "品目"), ("C", "数量"), ("D", "単価"), ("E", "金額")]
    for col, name in heads:
        c(f"{col}{header_row}", name, size=11, bold=True, color="FFFFFF",
          align="center", fill=navy_fill, border=box)
    ws.row_dimensions[header_row].height = 22

    first = header_row + 1
    # 明細行（最低8行の枠を確保して体裁を保つ）
    n_min = 8
    n = max(len(items), n_min)
    last = header_row + n
    for i in range(n):
        r = first + i
        c(f"A{r}", (i + 1) if i < len(items) else None, size=10, align="center", border=box)
        if i < len(items):
            name, qty, unit = items[i]
            c(f"B{r}", name, size=10, align="left", border=box, wrap=True)
            c(f"C{r}", qty, size=10, align="center", border=box, fmt='#,##0')
            c(f"D{r}", unit, size=10, align="right", border=box, fmt='"¥"#,##0')
            c(f"E{r}", f"=C{r}*D{r}", size=10, align="right", border=box, fmt='"¥"#,##0')
            ws.row_dimensions[r].height = 26
        else:
            for col in "BCDE":
                c(f"{col}{r}", None, size=10, border=box)
            ws.row_dimensions[r].height = 22

    # ---- 集計 ----
    sub_r = last + 1
    adj_r = last + 2
    tot_r = last + 3

    def summary(row, label, formula, bold=False, fill=None, label_color="000000"):
        ws.merge_cells(f"C{row}:D{row}")
        c(f"C{row}", label, size=11, bold=bold, align="right", color=label_color)
        c(f"E{row}", formula, size=11, bold=bold, align="right", border=box,
          fmt='"¥"#,##0', fill=fill)

    summary(sub_r, "小計", f"=SUM(E{first}:E{last})")
    summary(adj_r, "調整額（インボイス経過措置 8%）", f"=ROUND(E{sub_r}*0.08,0)")
    summary(tot_r, "合計（税込）", f"=E{sub_r}+E{adj_r}", bold=True,
            fill=band_fill, label_color=NAVY)
    ws.row_dimensions[tot_r].height = 26

    # ご請求金額（ヘッダ）に合計を参照
    c("C9", f"=E{tot_r}", size=16, bold=True, color=NAVY, align="center",
      fill=band_fill, fmt='"¥"#,##0')
    ws.row_dimensions[9].height = 30

    # ---- お振込先 ----
    bank_r = tot_r + 2
    c(f"A{bank_r}", "お振込先", size=11, bold=True, color=NAVY)
    ws[f"A{bank_r}"].border = Border(bottom=Side(style="thin", color=NAVY))
    for col in "BCDE":
        ws[f"{col}{bank_r}"].border = Border(bottom=Side(style="thin", color=NAVY))
    bank_lines = [
        "三井住友銀行　姫路支店",
        "普通　9554406",
        "口座名義：勝矢 夏帆（カツヤ カホ）",
    ]
    for i, line in enumerate(bank_lines):
        r = bank_r + 1 + i
        ws.merge_cells(f"A{r}:C{r}")
        c(f"A{r}", line, size=10)

    wb.save(path)
    # 検証用の合計計算
    sub = sum(q * u for _, q, u in items)
    adj = round(sub * 0.08)
    print(os.path.basename(path), "小計", sub, "調整額", adj, "合計", sub + adj)


july_items = [
    ("7/22 Westin横浜 対面ミーティング", 1, 5000),
    ("7/22 交通費（新幹線 新大阪→新横浜）", 1, 14190),
    ("7/22 交通費（現地移動 PASMO）", 1, 708),
    ("7/23 Six Senses Kyoto 交通費（新幹線 新横浜→京都）", 1, 13300),
    ("7/23 交通費（現地移動 PASMO）", 1, 1348),
]
august_items = [
    ("8/4 Aloft大阪 報酬", 1, 60000),
    ("8/4 交通費（バス 土佐堀2丁目→堂島）", 1, 210),
]

build(os.path.join(BASE, "invoice_soren_2026-07.xlsx"),
      "2026-07", "2026/07/31", "2026/08/31", july_items)
build(os.path.join(BASE, "invoice_soren_2026-08.xlsx"),
      "2026-08", "2026/08/31", "2026/09/30", august_items)

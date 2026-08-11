"""株式会社蒼蓮あて請求書（下書き）を、テンプレートに確定情報を差し込んで生成する。
交通費の金額・発行者情報・振込先は空欄（黄色）のまま。"""
import os
import openpyxl
from openpyxl.styles import Font, Alignment

BASE = os.path.dirname(os.path.abspath(__file__))
path = os.path.join(BASE, "invoice_soren_draft.xlsx")
wb = openpyxl.load_workbook(path)
ws = wb.active
FONT = "Arial"

def setv(coord, value, size=None, bold=None, color=None, align=None, wrap=None):
    c = ws[coord]
    c.value = value
    f = c.font
    c.font = Font(name=FONT,
                  size=size if size is not None else f.size,
                  bold=bold if bold is not None else f.bold,
                  color=color if color is not None else f.color)
    if align or wrap is not None:
        a = c.alignment
        c.alignment = Alignment(horizontal=align or a.horizontal,
                                vertical=a.vertical or "center",
                                wrap_text=wrap if wrap is not None else a.wrap_text)

# ---- 請求先 ----
setv("A5", "株式会社蒼蓮　御中")
setv("A6", "ご担当：小笹彩加 様")
setv("A7", "〒542-0012 大阪市中央区谷町6丁目3番25号 LOK09号")

# ---- メタ ----
setv("E4", "（例）2025-07")          # 請求書番号：任意
setv("E5", "")                        # 請求日：記入
setv("E6", "")                        # 支払期限：記入

# ---- 発行者（あなたの情報：記入してください）----
setv("A12", "（お名前 / 屋号を記入）")
setv("A13", "〒000-0000（ご住所を記入）")
setv("A14", "TEL / Email（連絡先を記入）")
setv("A15", "登録番号：（インボイス登録がある場合のみ T… を記入。未登録なら空欄）")

# ---- 明細 ----
# (row, 品目, 数量, 単価)  単価が None のセルは空欄（黄色）に残す
items = [
    (18, "7/22 Westin 対面ミーティング（報酬）", 1, 5000),
    (19, "7/22 交通費（実費）", 1, None),
    (20, "7/23 Six Senses Kyoto 交通費（研修・報酬なし）", 1, None),
    (21, "8/4 Aloft（報酬）", 1, 60000),
    (22, "8/4 交通費（実費）", 1, None),
    (23, "8/4 事務所までの交通費（8月以降対象）", 1, None),
]
for r in range(18, 26):
    ws[f"B{r}"].value = None
    ws[f"C{r}"].value = None
    ws[f"D{r}"].value = None
for r, name, qty, unit in items:
    setv(f"B{r}", name, size=10, align="left")
    ws[f"C{r}"].value = qty
    if unit is not None:
        ws[f"D{r}"].value = unit
    # D 空欄のときは黄色入力セルのまま（交通費を記入）

# ---- 集計ラベルを「調整額（経過措置8%）」に変更 ----
setv("C27", "調整額率（経過措置80%）", size=11, align="right")
ws["E27"].value = 0.08
setv("C28", "調整額（消費税相当）", size=11, align="right")
# 調整額 = 小計 × 8%（下記備考の前提。交通費実費を対象外にする場合は基準を調整してください）
ws["E28"].value = "=ROUND(E26*E27,0)"

# ---- 備考 ----
note = ("※ インボイス制度の経過措置（80%控除）に基づき、消費税10%ではなく調整額8%で計上しています。\n"
        "※ 交通費（実費）および事務所までの交通費の金額は記入欄（黄色）にご入力ください。\n"
        "※ 源泉徴収が必要な場合は「源泉徴収税額」欄に金額を入力すると合計から差し引かれます。\n"
        "※ ご提出：毎月5日までにメール（To: ayakaozasa@soren.co.jp / CC: ayakaanzo@soren.co.jp）")
# 備考セルを探す：A{note_r+1}。テンプレートでは合計行=30 → bank_r=32 → note_r=37 → 本文A38
setv("A38", note, size=9, wrap=True)

out = os.path.join(BASE, "invoice_soren_draft.xlsx")
wb.save(out)
print("saved", out)
# 検証
wb2 = openpyxl.load_workbook(out)
w = wb2.active
for c in ["A5", "E26", "C27", "E27", "C28", "E28", "E30", "B18", "D18", "B21", "D21"]:
    print(c, "->", repr(w[c].value))

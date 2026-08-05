import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "請求書"

# ---- Styles ----
FONT = "Arial"
input_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # yellow-ish = 入力
header_fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")  # navy
band_fill = PatternFill(start_color="EAEFF7", end_color="EAEFF7", fill_type="solid")

thin = Side(style="thin", color="B0B0B0")
med = Side(style="medium", color="203864")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

def cell(coord, value=None, size=11, bold=False, color="000000", align="left",
         fill=None, border=False, wrap=False, num_fmt=None, valign="center"):
    c = ws[coord]
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if fill:
        c.fill = fill
    if border:
        c.border = border_all
    if num_fmt:
        c.number_format = num_fmt
    return c

# ---- Column widths ----
widths = {"A": 5, "B": 30, "C": 12, "D": 14, "E": 16, "F": 3}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# ---- Title ----
ws.merge_cells("A1:F1")
cell("A1", "請求書", size=26, bold=True, color="203864", align="center")
ws.row_dimensions[1].height = 40

# ---- Legend ----
ws.merge_cells("A2:F2")
cell("A2", "※ 黄色のセルにご入力ください。金額・小計・消費税・合計は自動計算されます。",
     size=9, color="7F7F7F", align="left")

# ---- Meta (invoice no / dates) : right side ----
cell("D4", "請求書番号", size=10, bold=True, align="right")
cell("E4", "INV-0001", size=10, align="center", fill=input_fill, border=True)
cell("D5", "請求日", size=10, bold=True, align="right")
cell("E5", "2026/08/05", size=10, align="center", fill=input_fill, border=True, num_fmt="yyyy/mm/dd")
cell("D6", "支払期限", size=10, bold=True, align="right")
cell("E6", "2026/08/31", size=10, align="center", fill=input_fill, border=True, num_fmt="yyyy/mm/dd")

# ---- 請求先 (recipient) : left side ----
cell("A4", "請求先", size=11, bold=True, color="203864")
cell("A5", "株式会社サンプル", size=13, bold=True, fill=input_fill, border=True)
ws.merge_cells("A5:C5")
cell("A6", "御中 / ご担当者名", size=9, color="7F7F7F")
cell("A7", "〒000-0000 東京都〇〇区〇〇1-2-3", size=10, fill=input_fill, border=True)
ws.merge_cells("A7:C7")

# ---- ご請求金額 (headline total) ----
ws.merge_cells("A9:C9")
cell("A9", "ご請求金額（税込）", size=12, bold=True, color="FFFFFF", align="center", fill=header_fill)
ws.merge_cells("D9:E9")
cell("D9", "=E30", size=16, bold=True, align="center", color="203864",
     fill=band_fill, num_fmt='"¥"#,##0')
ws.row_dimensions[9].height = 28

# ---- 発行者 (issuer) block ----
cell("A11", "発行者", size=11, bold=True, color="203864")
issuer_rows = [
    ("A12", "〇〇〇〇（あなたの会社名・氏名）", True),
    ("A13", "〒000-0000 〇〇県〇〇市〇〇1-2-3", False),
    ("A14", "TEL: 000-0000-0000 / Email: sample@example.com", False),
    ("A15", "登録番号: T0000000000000", False),
]
for coord, val, bold in issuer_rows:
    r = coord[1:]
    ws.merge_cells(f"A{r}:E{r}")
    cell(coord, val, size=11 if bold else 9, bold=bold,
         color="000000" if bold else "595959", fill=input_fill, border=True)

# ---- 明細テーブル (line items) ----
tbl_top = 17
headers = [("A", "No."), ("B", "品目・摘要"), ("C", "数量"), ("D", "単価"), ("E", "金額")]
for col, name in headers:
    cell(f"{col}{tbl_top}", name, size=11, bold=True, color="FFFFFF",
         align="center", fill=header_fill, border=True)
ws.row_dimensions[tbl_top].height = 22

n_rows = 8
first = tbl_top + 1
last = tbl_top + n_rows
for i in range(n_rows):
    r = first + i
    cell(f"A{r}", i + 1, size=10, align="center", border=True)
    cell(f"B{r}", None, size=10, align="left", border=True, fill=input_fill)
    cell(f"C{r}", None, size=10, align="center", border=True, fill=input_fill, num_fmt='#,##0')
    cell(f"D{r}", None, size=10, align="right", border=True, fill=input_fill, num_fmt='"¥"#,##0')
    cell(f"E{r}", f"=IF(AND(C{r}<>\"\",D{r}<>\"\"),C{r}*D{r},\"\")",
         size=10, align="right", border=True, num_fmt='"¥"#,##0')
    ws.row_dimensions[r].height = 20

# Example row (first data row) — realistic sample, user overwrites
cell(f"B{first}", "Webサイト制作 一式", size=10, align="left", border=True, fill=input_fill)
cell(f"C{first}", 1, size=10, align="center", border=True, fill=input_fill, num_fmt='#,##0')
cell(f"D{first}", 200000, size=10, align="right", border=True, fill=input_fill, num_fmt='"¥"#,##0')

# ---- 集計 (subtotal / tax / total) ----
sub_r = last + 1        # 小計
rate_r = last + 2       # 消費税率
tax_r = last + 3        # 消費税
wh_r = last + 4         # 源泉徴収
tot_r = last + 5        # 合計 (=E30 referenced above)

def summary(row, label, formula, fmt='"¥"#,##0', input_cell=False, bold=False,
            label_color="000000", val_fill=None):
    cell(f"C{row}", label, size=11, bold=bold, align="right", color=label_color)
    ws.merge_cells(f"C{row}:D{row}")
    c = cell(f"E{row}", formula, size=11, bold=bold, align="right", border=True, num_fmt=fmt)
    if input_cell:
        c.fill = input_fill
    if val_fill:
        c.fill = val_fill
    return c

summary(sub_r, "小計", f"=SUM(E{first}:E{last})")
summary(rate_r, "消費税率", 0.10, fmt="0%", input_cell=True)
summary(tax_r, "消費税", f"=ROUND(E{sub_r}*E{rate_r},0)")
summary(wh_r, "源泉徴収税額（任意）", 0, input_cell=True)
summary(tot_r, "合計（税込）", f"=E{sub_r}+E{tax_r}-E{wh_r}", bold=True,
        label_color="203864", val_fill=band_fill)
ws.row_dimensions[tot_r].height = 24

# ---- 振込先 ----
bank_r = tot_r + 2
cell(f"A{bank_r}", "お振込先", size=11, bold=True, color="203864")
bank_lines = [
    "〇〇銀行 〇〇支店",
    "普通 0000000",
    "口座名義：〇〇〇〇（カ）",
]
for i, val in enumerate(bank_lines):
    r = bank_r + 1 + i
    ws.merge_cells(f"A{r}:E{r}")
    cell(f"A{r}", val, size=10, fill=input_fill, border=True)

# ---- 備考 ----
note_r = bank_r + len(bank_lines) + 2
cell(f"A{note_r}", "備考", size=11, bold=True, color="203864")
ws.merge_cells(f"A{note_r+1}:E{note_r+3}")
cell(f"A{note_r+1}", "（お支払いに関する連絡事項など）", size=10, fill=input_fill,
     border=True, wrap=True, valign="top")

ws.sheet_view.showGridLines = False

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "invoice_template.xlsx")
wb.save(out)
print("saved", out, "total_row=", tot_r)
